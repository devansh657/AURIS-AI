from __future__ import annotations

import hashlib
import json
import os
import re
import tomllib
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from tempfile import NamedTemporaryFile
from typing import Any, Callable

from auris.model_gateway import ModelReply, generate_reply
from auris.research_agent import run_research


DEFAULT_WORK_ROOT = Path.home() / "Documents" / "AURIS Work"
MAX_PROJECT_FILES = 16
MAX_FILE_BYTES = 80_000
MAX_PROJECT_BYTES = 320_000
ARTIFACT_MODEL_TIMEOUT_SECONDS = 12
ALLOWED_PROJECT_EXTENSIONS = {
    ".css", ".html", ".js", ".json", ".md", ".py", ".toml", ".txt", ".yaml", ".yml"
}
ALLOWED_SPECIAL_FILES = {".gitignore"}
HIGH_RISK_CODE_PATTERNS = (
    r"\bos\.system\s*\(",
    r"\bsubprocess\.",
    r"\beval\s*\(",
    r"\bexec\s*\(",
    r"\binvoke-expression\b",
    r"\bpowershell(?:\.exe)?\s+-enc",
    r"\brm\s+-rf\b",
    r"-----begin (?:rsa |ec |openssh )?private key-----",
)


@dataclass(frozen=True)
class WorkProductRequest:
    kind: str
    brief: str
    requested_name: str = ""
    output_format: str = ""


def match_work_product_command(command: str) -> WorkProductRequest | None:
    clean = _clean_command(command)
    if not clean:
        return None

    research_output = re.match(
        r"^(?:(?:conduct|do|perform)\s+)?(?:deep\s+|thorough\s+|comprehensive\s+)?"
        r"(?:research|investigate|look\s+up)\s+(.+?)"
        r"(?:\s+(?:and|then)\s+(?:create|write|generate|prepare|save|make)\s+"
        r"(?:me\s+)?(?:an?\s+)?(?:research\s+)?(?:(?:word|pdf)\s+)?(?:document|report|docx|pdf))$",
        clean,
        flags=re.IGNORECASE,
    )
    explicit_research_report = re.match(
        r"^(?:create|write|generate|prepare|make)\s+(?:me\s+)?(?:an?\s+)?research\s+report\s+"
        r"(?:about|on|for|covering)\s+(.+)$",
        clean,
        flags=re.IGNORECASE,
    )
    if research_output or explicit_research_report:
        match = research_output or explicit_research_report
        brief = match.group(1).strip(" .")
        return WorkProductRequest(
            "research_report",
            brief[:600],
            _requested_name(clean),
            "pdf" if "pdf" in clean.casefold() else "docx",
        )

    presentation = re.match(
        r"^(?:create|write|generate|prepare|draft|make|build)\s+(?:me\s+)?(?:an?\s+)?"
        r"(?:powerpoint\s+)?(?:presentation|slide\s+deck|slides|pptx)"
        r"(?:\s+(?:about|on|for|covering)\s+(.+))?$",
        clean,
        flags=re.IGNORECASE,
    )
    if presentation:
        subject = (presentation.group(1) or clean).strip(" .")
        return WorkProductRequest(
            "presentation", subject[:600], _requested_name(clean), "pptx"
        )

    spreadsheet = re.match(
        r"^(?:create|write|generate|prepare|draft|make|build)\s+(?:me\s+)?(?:an?\s+)?"
        r"(?:(?:microsoft|ms)\s+)?(?:excel\s+)?(?:spreadsheet|workbook|tracker|xlsx)"
        r"(?:\s+(?:about|on|for|covering|to\s+track)\s+(.+))?$",
        clean,
        flags=re.IGNORECASE,
    )
    if spreadsheet:
        subject = (spreadsheet.group(1) or clean).strip(" .")
        return WorkProductRequest(
            "spreadsheet", subject[:600], _requested_name(clean), "xlsx"
        )

    document = re.match(
        r"^(?:create|write|generate|prepare|draft|make)\s+(?:me\s+)?(?:an?\s+)?"
        r"((?:(?:microsoft|ms)\s+)?word\s+document|docx|pdf\s+(?:document|report)|pdf|"
        r"markdown\s+document|document|report|proposal|brief|essay)"
        r"(?:\s+(.+))?$",
        clean,
        flags=re.IGNORECASE,
    )
    if document:
        remainder = (document.group(2) or "").strip(" .")
        subject_match = re.search(
            r"\b(?:about|on|for|covering|that\s+covers)\s+(.+)$",
            remainder,
            flags=re.IGNORECASE,
        )
        subject = (
            subject_match.group(1)
            if subject_match
            else remainder if remainder and not re.match(r"^(?:called|named|titled)\b", remainder, flags=re.IGNORECASE)
            else clean
        ).strip(" .")
        requested_format = (
            "pdf" if "pdf" in document.group(1).casefold()
            else "md" if "markdown" in document.group(1).casefold()
            else "docx"
        )
        return WorkProductRequest(
            "document",
            subject[:600],
            _requested_name(clean),
            requested_format,
        )
    project = re.match(
        r"^(?:build|create|make|generate|scaffold|develop|code|program)\s+(?:me\s+)?(?:an?\s+)?(.+?\b"
        r"(?:code\s+project|software\s+project|python\s+project|web\s+app|application|app|website|"
        r"dashboard|api|chatbot|tool|project)\b.*)$",
        clean,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if project:
        return WorkProductRequest(
            "code_project",
            project.group(1).strip(" .")[:4_000],
            _requested_name(clean),
            "directory",
        )
    return None


def execute_work_product(
    request: WorkProductRequest,
    *,
    root: Path | None = None,
    model_generator: Callable[..., ModelReply] | None = None,
    research_runner: Callable[..., dict[str, Any]] | None = None,
) -> dict[str, Any]:
    work_root = _prepare_work_root(root or DEFAULT_WORK_ROOT)
    generator = model_generator or generate_reply
    if request.kind == "code_project":
        generative_projects = model_generator is not None or os.environ.get(
            "AURIS_GENERATIVE_PROJECTS", "0"
        ).strip() == "1"
        return _create_code_project(request, work_root, generator, generative=generative_projects)
    if request.kind == "document":
        return _create_document(request, work_root, generator)
    if request.kind == "presentation":
        return _create_presentation(request, work_root, generator)
    if request.kind == "spreadsheet":
        return _create_spreadsheet(request, work_root, generator)
    if request.kind == "research_report":
        return _create_research_report(
            request,
            work_root,
            research_runner or run_research,
        )
    return {"ok": False, "error": "That work-product type is not supported."}


def _create_code_project(
    request: WorkProductRequest,
    root: Path,
    generator: Callable[..., ModelReply],
    *,
    generative: bool,
) -> dict[str, Any]:
    reply: ModelReply | None = None
    generation_error = "latency-first verified scaffold selected"
    manifest = _fallback_project_manifest(request)
    if generative:
        try:
            reply = generator(
                [{"role": "user", "content": _project_prompt(request.brief)}],
                mode="artifact",
                project_name=None,
                project_instructions="",
                memories=[],
                latency_tier="fast",
                timeout_seconds=ARTIFACT_MODEL_TIMEOUT_SECONDS,
            )
            manifest = _project_manifest(reply.text)
            generation_error = ""
        except (RuntimeError, ValueError, TypeError, json.JSONDecodeError) as error:
            generation_error = str(error)[:500]
            manifest = _fallback_project_manifest(request)

    validation = _validate_project_manifest(manifest)
    if not validation["ok"]:
        generation_error = validation["error"]
        manifest = _fallback_project_manifest(request)
        validation = _validate_project_manifest(manifest)
    if not validation["ok"]:
        return {
            "ok": False,
            "error": validation["error"],
            "verification": "No project directory was created because the generated manifest failed validation.",
        }

    title = request.requested_name or str(manifest.get("name") or "AURIS Project")
    project_root = _create_unique_directory(root / "Projects", _slug(title, "auris-project"))
    written: list[dict[str, Any]] = []
    try:
        for file_spec in manifest["files"]:
            relative = PurePosixPath(file_spec["path"])
            target = project_root.joinpath(*relative.parts)
            target.parent.mkdir(parents=True, exist_ok=True)
            content = file_spec["content"]
            with target.open("x", encoding="utf-8", newline="") as handle:
                handle.write(content)
            encoded = content.encode("utf-8")
            written.append(
                {
                    "path": relative.as_posix(),
                    "bytes": len(encoded),
                    "sha256": hashlib.sha256(encoded).hexdigest(),
                }
            )
    except OSError as error:
        return {
            "ok": False,
            "error": f"The managed project write failed: {error}",
            "artifact": {"path": str(project_root), "files": written, "incomplete": True},
            "verification": "The incomplete project remains visible for inspection; AURIS did not report it as complete.",
        }

    checks = _verify_project(project_root, written)
    manifest_record = {
        "schema": "auris.work-product.v1",
        "kind": "code_project",
        "brief": request.brief,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "generator": {
            "provider": getattr(reply, "provider", "auris_template_engine"),
            "model": getattr(reply, "model", "verified_scaffold_v1"),
            "route": getattr(reply, "route", "fallback"),
            "duration_ms": getattr(reply, "duration_ms", 0),
            "fallback_reason": generation_error or None,
        },
        "files": written,
        "checks": checks,
        "generated_code_executed": False,
    }
    manifest_path = project_root / "AURIS-MANIFEST.json"
    with manifest_path.open("x", encoding="utf-8", newline="") as handle:
        json.dump(manifest_record, handle, ensure_ascii=True, indent=2)
        handle.write("\n")
    manifest_hash = _sha256_file(manifest_path)
    ok = all(check["passed"] for check in checks)
    return {
        "ok": ok,
        "complete": ok,
        "operation": "create_code_project",
        "message": (
            f"Created and statically verified {title} with {len(written)} project files at {project_root}. "
            "Generated code was not executed; launch instructions are in the project README."
            if ok
            else f"Created {title}, but one or more static checks failed. I did not execute the generated code."
        ),
        "verification": (
            f"Verified: {len(written)} files were written without overwrite, hashes were recorded, and all {len(checks)} static checks passed. Generated code was not executed."
            if ok
            else "The project files and hashes were preserved, but static verification did not pass. Generated code was not executed."
        ),
        "artifact": {
            "kind": "code_project",
            "name": title,
            "path": str(project_root),
            "manifest_path": str(manifest_path),
            "manifest_sha256": manifest_hash,
            "files": written,
            "checks": checks,
            "generated_code_executed": False,
        },
        "generation": manifest_record["generator"],
    }


def _create_document(
    request: WorkProductRequest,
    root: Path,
    generator: Callable[..., ModelReply],
) -> dict[str, Any]:
    title = request.requested_name or _title_from_brief(request.brief, "AURIS Document")
    try:
        reply = generator(
            [{"role": "user", "content": _document_prompt(title, request.brief)}],
            mode="artifact",
            project_name=None,
            project_instructions="",
            memories=[],
            latency_tier="fast",
            timeout_seconds=ARTIFACT_MODEL_TIMEOUT_SECONDS,
        )
        content = _document_structure(reply.text, title, request.brief)
        generator_info = {"provider": reply.provider, "model": reply.model, "route": reply.route, "duration_ms": reply.duration_ms}
    except RuntimeError as error:
        content = _document_structure("", title, request.brief)
        generator_info = {
            "provider": "deterministic_fallback",
            "model": "built_in_document",
            "fallback_reason": str(error)[:500],
        }
    if request.output_format == "pdf":
        return _write_pdf(content, root / "Documents", title, "document", generator_info)
    if request.output_format == "md":
        return _write_markdown(content, root / "Documents", title, "document", generator_info)
    return _write_docx(content, root / "Documents", title, "document", generator_info)


def _create_research_report(
    request: WorkProductRequest,
    root: Path,
    research_runner: Callable[..., dict[str, Any]],
) -> dict[str, Any]:
    research = research_runner(f"research {request.brief}")
    if not research.get("ok"):
        return {
            "ok": False,
            "error": research.get("error", "Research did not complete."),
            "verification": research.get("verification", "No report was created because research evidence was unavailable."),
            "research_action": research,
        }
    title = request.requested_name or _title_from_brief(request.brief, "Research Report")
    content = _research_document_structure(title, request.brief, research)
    writer = _write_pdf if request.output_format == "pdf" else _write_docx
    result = writer(
        content, root / "Research", title, "research_report",
        {"provider": "auris_research_engine", "model": "evidence_ledger"},
    )
    result["research_action"] = research
    result["complete"] = bool(result.get("ok") and research.get("definition_of_done_met"))
    result["message"] = (
        f"Created the evidence-backed research report at {result['artifact']['path']}."
        if result.get("ok")
        else result.get("message", "The research report could not be created.")
    )
    result["verification"] = (
        result.get("verification", "")
        + " Research status: "
        + research.get("verification", "No research verification was supplied.")
    ).strip()
    return result


def _write_docx(
    content: dict[str, Any],
    directory: Path,
    title: str,
    kind: str,
    generator_info: dict[str, Any],
) -> dict[str, Any]:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError:
        return {
            "ok": False,
            "error": "The local Word document writer is unavailable.",
            "verification": "No document file was created.",
        }

    directory.mkdir(parents=True, exist_ok=True)
    if directory.is_symlink():
        return {"ok": False, "error": "The managed document directory cannot be a link."}
    path = _unique_path(directory, _slug(title, "auris-document"), ".docx")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    document.core_properties.title = title
    document.core_properties.author = "AURIS"
    document.add_heading(content["title"], level=0)
    if content.get("summary"):
        paragraph = document.add_paragraph(content["summary"])
        paragraph.style = document.styles["Subtitle"]
    for section_data in content.get("sections", []):
        document.add_heading(section_data["heading"], level=1)
        for paragraph_text in section_data.get("paragraphs", []):
            document.add_paragraph(paragraph_text)
        for item in section_data.get("bullets", []):
            document.add_paragraph(item, style="List Bullet")
    with NamedTemporaryFile(suffix=".docx", delete=False, dir=directory) as temporary:
        temporary_path = Path(temporary.name)
    try:
        document.save(temporary_path)
        payload = temporary_path.read_bytes()
        with path.open("xb") as target:
            target.write(payload)
    finally:
        temporary_path.unlink(missing_ok=True)

    verification = _verify_docx(path)
    artifact = {
        "kind": kind,
        "name": title,
        "path": str(path),
        "format": "docx",
        "bytes": path.stat().st_size,
        "sha256": _sha256_file(path),
        "paragraphs": verification.get("paragraphs", 0),
        "generator": generator_info,
        "checks": verification.get("checks", []),
    }
    ok = bool(verification.get("ok"))
    return {
        "ok": ok,
        "complete": ok,
        "operation": f"create_{kind}",
        "message": (
            f"Created and verified the Word document {path.name} at {path}."
            if ok
            else f"Created {path.name}, but its Word-package verification failed."
        ),
        "verification": verification["verification"],
        "artifact": artifact,
    }


def _create_presentation(
    request: WorkProductRequest,
    root: Path,
    generator: Callable[..., ModelReply],
) -> dict[str, Any]:
    title = request.requested_name or _title_from_brief(request.brief, "AURIS Presentation")
    try:
        reply = generator(
            [{"role": "user", "content": _presentation_prompt(title, request.brief)}],
            mode="artifact",
            project_name=None,
            project_instructions="",
            memories=[],
            latency_tier="fast",
            timeout_seconds=ARTIFACT_MODEL_TIMEOUT_SECONDS,
        )
        content = _presentation_structure(reply.text, title, request.brief)
        generator_info = {"provider": reply.provider, "model": reply.model, "route": reply.route, "duration_ms": reply.duration_ms}
    except RuntimeError as error:
        content = _presentation_structure("", title, request.brief)
        generator_info = {
            "provider": "deterministic_fallback",
            "model": "built_in_presentation",
            "fallback_reason": str(error)[:500],
        }
    return _write_pptx(content, root / "Presentations", title, generator_info)


def _create_spreadsheet(
    request: WorkProductRequest,
    root: Path,
    generator: Callable[..., ModelReply],
) -> dict[str, Any]:
    title = request.requested_name or _title_from_brief(request.brief, "AURIS Workbook")
    try:
        reply = generator(
            [{"role": "user", "content": _spreadsheet_prompt(title, request.brief)}],
            mode="artifact",
            project_name=None,
            project_instructions="",
            memories=[],
            latency_tier="fast",
            timeout_seconds=ARTIFACT_MODEL_TIMEOUT_SECONDS,
        )
        content = _spreadsheet_structure(reply.text, title, request.brief)
        generator_info = {"provider": reply.provider, "model": reply.model, "route": reply.route, "duration_ms": reply.duration_ms}
    except RuntimeError as error:
        content = _spreadsheet_structure("", title, request.brief)
        generator_info = {
            "provider": "deterministic_fallback",
            "model": "built_in_workbook",
            "fallback_reason": str(error)[:500],
        }
    return _write_xlsx(content, root / "Spreadsheets", title, generator_info)


def _write_pdf(
    content: dict[str, Any],
    directory: Path,
    title: str,
    kind: str,
    generator_info: dict[str, Any],
) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        return {"ok": False, "error": "The local PDF writer is unavailable.", "verification": "No PDF was created."}

    directory.mkdir(parents=True, exist_ok=True)
    path = _unique_path(directory, _slug(title, "auris-document"), ".pdf")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="AurisTitle", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=8 * mm))
    story: list[Any] = [Paragraph(_xml_text(content["title"]), styles["AurisTitle"])]
    if content.get("summary"):
        story.extend([Paragraph(_xml_text(content["summary"]), styles["BodyText"]), Spacer(1, 5 * mm)])
    for section_data in content.get("sections", []):
        story.append(Paragraph(_xml_text(section_data["heading"]), styles["Heading1"]))
        for paragraph in section_data.get("paragraphs", []):
            story.extend([Paragraph(_xml_text(paragraph), styles["BodyText"]), Spacer(1, 2 * mm)])
        bullets = section_data.get("bullets", [])
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(_xml_text(item), styles["BodyText"])) for item in bullets], bulletType="bullet"))
            story.append(Spacer(1, 3 * mm))
    with NamedTemporaryFile(suffix=".pdf", delete=False, dir=directory) as temporary:
        temporary_path = Path(temporary.name)
    try:
        SimpleDocTemplate(str(temporary_path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm).build(story)
        payload = temporary_path.read_bytes()
        with path.open("xb") as target:
            target.write(payload)
    finally:
        temporary_path.unlink(missing_ok=True)
    try:
        reader = PdfReader(str(path))
        extracted = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        checks = [
            {"name": "pdf_reopen", "passed": bool(reader.pages), "evidence": f"{len(reader.pages)} pages reopened"},
            {"name": "pdf_text", "passed": len(extracted) >= 20, "evidence": f"{len(extracted)} extracted characters"},
        ]
    except (OSError, ValueError) as error:
        checks = [{"name": "pdf_reopen", "passed": False, "evidence": str(error)[:300]}]
    ok = all(check["passed"] for check in checks)
    artifact = _file_artifact(kind, title, path, "pdf", generator_info, checks)
    artifact["pages"] = len(reader.pages) if ok else 0
    return {
        "ok": ok,
        "complete": ok,
        "operation": f"create_{kind}",
        "message": f"Created and verified the PDF {path.name} at {path}." if ok else f"Created {path.name}, but PDF verification failed.",
        "verification": f"Verified: the PDF reopened with {artifact['pages']} pages and extractable text; size and SHA-256 were recorded." if ok else "The PDF exists, but it did not pass reopen and text checks.",
        "artifact": artifact,
    }


def _write_markdown(
    content: dict[str, Any],
    directory: Path,
    title: str,
    kind: str,
    generator_info: dict[str, Any],
) -> dict[str, Any]:
    directory.mkdir(parents=True, exist_ok=True)
    path = _unique_path(directory, _slug(title, "auris-document"), ".md")
    lines = [f"# {content['title']}", "", content.get("summary", "").strip(), ""]
    for section in content.get("sections", []):
        lines.extend([f"## {section['heading']}", ""])
        for paragraph in section.get("paragraphs", []):
            lines.extend([paragraph, ""])
        for bullet in section.get("bullets", []):
            lines.append(f"- {bullet}")
        lines.append("")
    text = "\n".join(lines).strip() + "\n"
    with path.open("x", encoding="utf-8", newline="") as handle:
        handle.write(text)
    reopened = path.read_text(encoding="utf-8")
    checks = [{"name": "markdown_reopen", "passed": reopened == text, "evidence": f"{len(reopened)} characters reopened"}]
    artifact = _file_artifact(kind, title, path, "md", generator_info, checks)
    return {
        "ok": True,
        "complete": True,
        "operation": f"create_{kind}",
        "message": f"Created and verified the Markdown document {path.name} at {path}.",
        "verification": "Verified: the UTF-8 Markdown content reopened byte-for-byte and its SHA-256 was recorded.",
        "artifact": artifact,
    }


def _write_pptx(
    content: dict[str, Any],
    directory: Path,
    title: str,
    generator_info: dict[str, Any],
) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return {"ok": False, "error": "The local PowerPoint writer is unavailable.", "verification": "No presentation was created."}
    directory.mkdir(parents=True, exist_ok=True)
    path = _unique_path(directory, _slug(title, "auris-presentation"), ".pptx")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = content["title"]
    title_slide.placeholders[1].text = content.get("subtitle", "Prepared by AURIS")
    for slide_data in content.get("slides", [])[:20]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_data["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(slide_data.get("bullets", [])[:8]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(22)
    with NamedTemporaryFile(suffix=".pptx", delete=False, dir=directory) as temporary:
        temporary_path = Path(temporary.name)
    try:
        presentation.save(temporary_path)
        payload = temporary_path.read_bytes()
        with path.open("xb") as target:
            target.write(payload)
    finally:
        temporary_path.unlink(missing_ok=True)
    try:
        reopened = Presentation(path)
        slide_count = len(reopened.slides)
        titled = sum(bool(slide.shapes.title and slide.shapes.title.text.strip()) for slide in reopened.slides)
        checks = [
            {"name": "pptx_reopen", "passed": slide_count >= 2, "evidence": f"{slide_count} slides reopened"},
            {"name": "pptx_titles", "passed": titled == slide_count, "evidence": f"{titled}/{slide_count} slides titled"},
        ]
    except (OSError, ValueError, KeyError) as error:
        checks = [{"name": "pptx_reopen", "passed": False, "evidence": str(error)[:300]}]
        slide_count = 0
    ok = all(check["passed"] for check in checks)
    artifact = _file_artifact("presentation", title, path, "pptx", generator_info, checks)
    artifact["slides"] = slide_count
    return {
        "ok": ok,
        "complete": ok,
        "operation": "create_presentation",
        "message": f"Created and verified the PowerPoint deck {path.name} at {path}." if ok else f"Created {path.name}, but presentation verification failed.",
        "verification": f"Verified: the PPTX reopened with {slide_count} titled slides; size and SHA-256 were recorded." if ok else "The PPTX exists, but it did not pass reopen and slide checks.",
        "artifact": artifact,
    }


def _write_xlsx(
    content: dict[str, Any],
    directory: Path,
    title: str,
    generator_info: dict[str, Any],
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return {"ok": False, "error": "The local Excel writer is unavailable.", "verification": "No workbook was created."}
    directory.mkdir(parents=True, exist_ok=True)
    path = _unique_path(directory, _slug(title, "auris-workbook"), ".xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AURIS Data"
    headers = content["headers"]
    sheet.append(headers)
    for row in content.get("rows", [])[:200]:
        sheet.append([_safe_spreadsheet_value(value) for value in row[: len(headers)]])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087E8B")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    with NamedTemporaryFile(suffix=".xlsx", delete=False, dir=directory) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        payload = temporary_path.read_bytes()
        with path.open("xb") as target:
            target.write(payload)
    finally:
        temporary_path.unlink(missing_ok=True)
    try:
        reopened = load_workbook(path, read_only=True, data_only=False)
        reopened_sheet = reopened["AURIS Data"]
        rows = reopened_sheet.max_row
        columns = reopened_sheet.max_column
        formulas = sum(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for row in reopened_sheet.iter_rows()
            for cell in row
        )
        reopened.close()
        checks = [
            {"name": "xlsx_reopen", "passed": rows >= 1 and columns == len(headers), "evidence": f"{rows} rows x {columns} columns reopened"},
            {"name": "xlsx_formula_safety", "passed": formulas == 0, "evidence": f"{formulas} formula cells"},
        ]
    except (OSError, ValueError, KeyError) as error:
        checks = [{"name": "xlsx_reopen", "passed": False, "evidence": str(error)[:300]}]
        rows = columns = 0
    ok = all(check["passed"] for check in checks)
    artifact = _file_artifact("spreadsheet", title, path, "xlsx", generator_info, checks)
    artifact.update({"rows": rows, "columns": columns})
    return {
        "ok": ok,
        "complete": ok,
        "operation": "create_spreadsheet",
        "message": f"Created and verified the Excel workbook {path.name} at {path}." if ok else f"Created {path.name}, but workbook verification failed.",
        "verification": f"Verified: the XLSX reopened as {rows} rows by {columns} columns with no active formulas; size and SHA-256 were recorded." if ok else "The XLSX exists, but it did not pass reopen and formula-safety checks.",
        "artifact": artifact,
    }


def _project_prompt(brief: str) -> str:
    return (
        "Create a small, coherent, runnable starter project for this request: "
        f"{brief}. Return only strict JSON with keys name, summary, and files. files must be an array "
        "of objects with path and content. Use at most 10 UTF-8 text files and no binary data, shell "
        "scripts, secrets, subprocesses, eval, or destructive operations. Include README.md with exact "
        "local run instructions. Prefer standard-library or browser-native code so the project works "
        "without installing dependencies."
    )


def _document_prompt(title: str, brief: str) -> str:
    return (
        f"Draft a useful professional document titled {title!r} about {brief!r}. Return only strict JSON "
        "with title, summary, and sections. sections is an array of objects with heading, paragraphs "
        "(array of prose strings), and bullets (array of concise strings). Do not invent sources, metrics, "
        "people, dates, or completed actions."
    )


def _presentation_prompt(title: str, brief: str) -> str:
    return (
        f"Create a concise professional presentation titled {title!r} about {brief!r}. Return only strict "
        "JSON with title, subtitle, and slides. slides is an array of 4 to 10 objects with title and "
        "bullets (2 to 6 concise strings). Do not invent sources, metrics, people, or completed actions."
    )


def _spreadsheet_prompt(title: str, brief: str) -> str:
    return (
        f"Design a useful spreadsheet titled {title!r} for {brief!r}. Return only strict JSON with title, "
        "headers (2 to 12 concise column names), and rows (up to 30 arrays). Use only plain text or numeric "
        "cell values. Do not return formulas, links, macros, personal data, or invented factual records. "
        "Use clearly labelled example rows only when the user supplied no real data."
    )


def _project_manifest(text: str) -> dict[str, Any]:
    payload = _json_object(text)
    if not isinstance(payload, dict):
        raise ValueError("The model did not return a project object.")
    return payload


def _validate_project_manifest(manifest: dict[str, Any]) -> dict[str, Any]:
    files = manifest.get("files")
    if not isinstance(files, list) or not 1 <= len(files) <= MAX_PROJECT_FILES:
        return {"ok": False, "error": f"A project must contain 1 to {MAX_PROJECT_FILES} text files."}
    seen: set[str] = set()
    total = 0
    for item in files:
        if not isinstance(item, dict) or set(item) != {"path", "content"}:
            return {"ok": False, "error": "Every project file must contain only path and content."}
        path_text = str(item["path"]).replace("\\", "/").strip()
        content = item["content"]
        path = PurePosixPath(path_text)
        if (
            not path_text
            or path.is_absolute()
            or ".." in path.parts
            or any(part in {"", "."} or part.startswith(".") for part in path.parts[:-1])
            or path.name.startswith(".") and path.name not in ALLOWED_SPECIAL_FILES
            or (path.suffix.casefold() not in ALLOWED_PROJECT_EXTENSIONS and path.name not in ALLOWED_SPECIAL_FILES)
        ):
            return {"ok": False, "error": f"Unsafe or unsupported project path: {path_text[:120]}"}
        normalised = path.as_posix().casefold()
        if normalised in seen:
            return {"ok": False, "error": f"Duplicate project path: {path_text[:120]}"}
        if not isinstance(content, str):
            return {"ok": False, "error": f"Project content for {path_text[:120]} is not text."}
        size = len(content.encode("utf-8"))
        if size > MAX_FILE_BYTES:
            return {"ok": False, "error": f"Project file exceeds {MAX_FILE_BYTES} bytes: {path_text[:120]}"}
        if any(re.search(pattern, content, flags=re.IGNORECASE) for pattern in HIGH_RISK_CODE_PATTERNS):
            return {"ok": False, "error": f"High-risk code was rejected in {path_text[:120]}."}
        seen.add(normalised)
        total += size
    if total > MAX_PROJECT_BYTES:
        return {"ok": False, "error": f"Project exceeds the {MAX_PROJECT_BYTES}-byte generation limit."}
    return {"ok": True, "error": ""}


def _fallback_project_manifest(request: WorkProductRequest) -> dict[str, Any]:
    brief = request.brief
    title = request.requested_name or _title_from_brief(brief, "AURIS Project")
    lowered = brief.casefold()
    if "chatbot" in lowered or "assistant" in lowered:
        return _fallback_chatbot_project(title, brief)
    if re.search(r"\bapi\b", lowered):
        return _fallback_api_project(title, brief)
    if any(term in lowered for term in ("website", "web app", "dashboard", "frontend")):
        return _fallback_web_project(title, brief)
    return _fallback_python_project(title, brief)


def _fallback_chatbot_project(title: str, brief: str) -> dict[str, Any]:
    return {
        "name": title,
        "summary": brief,
        "files": [
            {
                "path": "README.md",
                "content": (
                    f"# {title}\n\n{brief}\n\nThis CLI chatbot uses the local Ollama API and Python's standard library. "
                    "Start Ollama, ensure `gemma3:4b` is available, then run `python app.py`.\n"
                ),
            },
            {
                "path": "app.py",
                "content": (
                    "from __future__ import annotations\n\nimport json\nimport urllib.request\n\n"
                    "ENDPOINT = 'http://127.0.0.1:11434/api/chat'\nMODEL = 'gemma3:4b'\n\n"
                    "def reply(history: list[dict[str, str]]) -> str:\n"
                    "    payload = json.dumps({'model': MODEL, 'messages': history, 'stream': False}).encode('utf-8')\n"
                    "    request = urllib.request.Request(ENDPOINT, data=payload, headers={'Content-Type': 'application/json'})\n"
                    "    with urllib.request.urlopen(request, timeout=90) as response:\n"
                    "        result = json.loads(response.read().decode('utf-8'))\n"
                    "    return str(result['message']['content']).strip()\n\n"
                    "def main() -> None:\n"
                    "    history: list[dict[str, str]] = []\n"
                    "    print('Chatbot ready. Type exit to stop.')\n"
                    "    while True:\n"
                    "        text = input('You: ').strip()\n"
                    "        if text.casefold() in {'exit', 'quit'}:\n"
                    "            break\n"
                    "        if not text:\n"
                    "            continue\n"
                    "        history.append({'role': 'user', 'content': text})\n"
                    "        answer = reply(history)\n"
                    "        history.append({'role': 'assistant', 'content': answer})\n"
                    "        print(f'Assistant: {answer}')\n\n"
                    "if __name__ == '__main__':\n"
                    "    main()\n"
                ),
            },
            {
                "path": "pyproject.toml",
                "content": "[project]\nname = \"auris-generated-chatbot\"\nversion = \"0.1.0\"\nrequires-python = \">=3.11\"\n",
            },
        ],
    }


def _fallback_api_project(title: str, brief: str) -> dict[str, Any]:
    return {
        "name": title,
        "summary": brief,
        "files": [
            {
                "path": "README.md",
                "content": f"# {title}\n\n{brief}\n\nRun `python server.py`, then open `http://127.0.0.1:8080/health`.\n",
            },
            {
                "path": "server.py",
                "content": (
                    "from __future__ import annotations\n\nimport json\nfrom http.server import BaseHTTPRequestHandler, ThreadingHTTPServer\n\n"
                    "class Handler(BaseHTTPRequestHandler):\n"
                    "    def do_GET(self) -> None:\n"
                    "        status = 200 if self.path == '/health' else 404\n"
                    "        body = {'ok': True, 'service': 'generated-api'} if status == 200 else {'ok': False, 'error': 'not found'}\n"
                    "        payload = json.dumps(body).encode('utf-8')\n"
                    "        self.send_response(status)\n"
                    "        self.send_header('Content-Type', 'application/json')\n"
                    "        self.send_header('Content-Length', str(len(payload)))\n"
                    "        self.end_headers()\n"
                    "        self.wfile.write(payload)\n\n"
                    "if __name__ == '__main__':\n"
                    "    ThreadingHTTPServer(('127.0.0.1', 8080), Handler).serve_forever()\n"
                ),
            },
            {
                "path": "pyproject.toml",
                "content": "[project]\nname = \"auris-generated-api\"\nversion = \"0.1.0\"\nrequires-python = \">=3.11\"\n",
            },
        ],
    }


def _fallback_web_project(title: str, brief: str) -> dict[str, Any]:
    safe_title = _html_text(title)
    safe_brief = _html_text(brief)
    return {
        "name": title,
        "summary": brief,
        "files": [
            {
                "path": "README.md",
                "content": f"# {title}\n\n{brief}\n\nOpen `index.html` in a browser. No build step is required.\n",
            },
            {
                "path": "index.html",
                "content": (
                    "<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
                    "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">\n"
                    f"<title>{safe_title}</title>\n<link rel=\"stylesheet\" href=\"styles.css\">\n</head>\n"
                    f"<body><main><h1>{safe_title}</h1><p>{safe_brief}</p><form id=\"task-form\">"
                    "<label for=\"task\">New item</label><input id=\"task\" required maxlength=\"120\">"
                    "<button>Add</button></form><ul id=\"items\"></ul></main><script src=\"app.js\"></script></body></html>\n"
                ),
            },
            {
                "path": "styles.css",
                "content": (
                    ":root{font-family:system-ui,sans-serif;color:#10212b;background:#eef4f5}"
                    "body{margin:0}main{max-width:760px;margin:48px auto;padding:28px;background:#fff;border:1px solid #b9cbd0}"
                    "form{display:flex;gap:8px}input{flex:1;padding:10px}button{padding:10px 16px;background:#087e8b;color:white;border:0}"
                    "li{display:flex;justify-content:space-between;padding:12px 0;border-bottom:1px solid #d7e2e5}\n"
                ),
            },
            {
                "path": "app.js",
                "content": (
                    "const form=document.querySelector('#task-form');const input=document.querySelector('#task');"
                    "const list=document.querySelector('#items');form.addEventListener('submit',(event)=>{event.preventDefault();"
                    "const item=document.createElement('li');const text=document.createElement('span');text.textContent=input.value;"
                    "const remove=document.createElement('button');remove.type='button';remove.textContent='Remove';"
                    "remove.addEventListener('click',()=>item.remove());item.append(text,remove);list.append(item);input.value='';input.focus();});\n"
                ),
            },
        ],
    }


def _fallback_python_project(title: str, brief: str) -> dict[str, Any]:
    description = brief.replace('"', "'")
    return {
        "name": title,
        "summary": brief,
        "files": [
            {
                "path": "README.md",
                "content": f"# {title}\n\n{brief}\n\nRun with `python main.py`. The project uses only the Python standard library.\n",
            },
            {
                "path": "main.py",
                "content": (
                    "from __future__ import annotations\n\n"
                    f"PROJECT_DESCRIPTION = {description!r}\n\n"
                    "def main() -> None:\n"
                    "    print(PROJECT_DESCRIPTION)\n"
                    "    value = input('Enter a task (or press Enter to exit): ').strip()\n"
                    "    if value:\n"
                    "        print(f'Recorded: {value}')\n\n"
                    "if __name__ == '__main__':\n"
                    "    main()\n"
                ),
            },
            {
                "path": "pyproject.toml",
                "content": (
                    "[project]\nname = \"auris-generated-project\"\nversion = \"0.1.0\"\n"
                    f"description = {json.dumps(brief[:200])}\nrequires-python = \">=3.11\"\n"
                ),
            },
        ],
    }


def _verify_project(root: Path, written: list[dict[str, Any]]) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    for record in written:
        path = root.joinpath(*PurePosixPath(record["path"]).parts)
        digest = _sha256_file(path) if path.is_file() else ""
        checks.append(
            {
                "name": f"hash:{record['path']}",
                "passed": digest == record["sha256"],
                "evidence": digest or "file missing",
            }
        )
        try:
            if path.suffix.casefold() == ".py":
                compile(path.read_text(encoding="utf-8"), str(path), "exec")
            elif path.suffix.casefold() == ".json":
                json.loads(path.read_text(encoding="utf-8"))
            elif path.name == "pyproject.toml" or path.suffix.casefold() == ".toml":
                tomllib.loads(path.read_text(encoding="utf-8"))
            elif path.suffix.casefold() == ".html":
                html = path.read_text(encoding="utf-8").casefold()
                if "<html" not in html or "</html>" not in html:
                    raise ValueError("missing HTML root element")
        except (SyntaxError, ValueError, tomllib.TOMLDecodeError, UnicodeError) as error:
            checks.append(
                {"name": f"syntax:{record['path']}", "passed": False, "evidence": str(error)[:300]}
            )
        else:
            if path.suffix.casefold() in {".py", ".json", ".toml", ".html"}:
                checks.append(
                    {"name": f"syntax:{record['path']}", "passed": True, "evidence": "static parser accepted file"}
                )
    return checks


def _document_structure(text: str, title: str, brief: str) -> dict[str, Any]:
    try:
        payload = _json_object(text)
    except (ValueError, json.JSONDecodeError):
        payload = {}
    if isinstance(payload, dict) and isinstance(payload.get("sections"), list):
        sections = _safe_sections(payload["sections"])
        if sections:
            return {
                "title": str(payload.get("title") or title)[:200],
                "summary": str(payload.get("summary") or "")[:1500],
                "sections": sections,
            }
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    if paragraphs:
        sections = [{"heading": "Overview", "paragraphs": paragraphs[1:8] or [paragraphs[0]], "bullets": []}]
        summary = paragraphs[0][:1500]
    else:
        summary = f"Structured working document for: {brief}"
        sections = [
            {"heading": "Objective", "paragraphs": [f"Define and deliver the requested outcome for {brief}."], "bullets": []},
            {"heading": "Scope", "paragraphs": ["Confirm the intended audience, constraints, inputs, and required level of detail before consequential use."], "bullets": []},
            {"heading": "Key Considerations", "paragraphs": [], "bullets": ["Use verified source material", "Separate evidence from assumptions", "Record unresolved questions", "Review the final output for accuracy"]},
            {"heading": "Next Steps", "paragraphs": ["Replace provisional content with verified project-specific evidence, assign ownership, and record completion criteria."], "bullets": []},
        ]
    return {"title": title, "summary": summary, "sections": sections}


def _presentation_structure(text: str, title: str, brief: str) -> dict[str, Any]:
    try:
        payload = _json_object(text)
    except (ValueError, json.JSONDecodeError):
        payload = {}
    slides: list[dict[str, Any]] = []
    if isinstance(payload, dict) and isinstance(payload.get("slides"), list):
        for item in payload["slides"][:20]:
            if not isinstance(item, dict):
                continue
            slide_title = str(item.get("title") or "").strip()[:160]
            bullets = [str(value).strip()[:800] for value in item.get("bullets", [])[:8] if str(value).strip()]
            if slide_title and bullets:
                slides.append({"title": slide_title, "bullets": bullets})
    if not slides:
        slides = [
            {"title": "Objective", "bullets": [brief, "Define the intended audience and outcome"]},
            {"title": "Key Considerations", "bullets": ["Current context", "Constraints and dependencies", "Evidence required"]},
            {"title": "Recommended Approach", "bullets": ["Prioritise the highest-value work", "Validate assumptions before execution", "Track measurable outcomes"]},
            {"title": "Next Steps", "bullets": ["Confirm scope", "Assign owners and dates", "Review progress and evidence"]},
        ]
    return {
        "title": str(payload.get("title") or title)[:200] if isinstance(payload, dict) else title,
        "subtitle": str(payload.get("subtitle") or f"Prepared for: {brief}")[:300] if isinstance(payload, dict) else f"Prepared for: {brief}",
        "slides": slides,
    }


def _spreadsheet_structure(text: str, title: str, brief: str) -> dict[str, Any]:
    try:
        payload = _json_object(text)
    except (ValueError, json.JSONDecodeError):
        payload = {}
    headers: list[str] = []
    rows: list[list[Any]] = []
    if isinstance(payload, dict) and isinstance(payload.get("headers"), list):
        headers = [str(value).strip()[:100] for value in payload["headers"][:12] if str(value).strip()]
        raw_rows = payload.get("rows") if isinstance(payload.get("rows"), list) else []
        for raw_row in raw_rows[:200]:
            if isinstance(raw_row, list):
                rows.append([_safe_spreadsheet_value(value) for value in raw_row[: len(headers)]])
    if len(headers) < 2:
        headers = ["Item", "Status", "Owner", "Due Date", "Notes"]
        rows = [[f"Example: {brief}", "Not started", "", "", "Replace this example with verified data"]]
    return {"title": str(payload.get("title") or title)[:200] if isinstance(payload, dict) else title, "headers": headers, "rows": rows}


def _research_document_structure(
    title: str, brief: str, research: dict[str, Any]
) -> dict[str, Any]:
    claims = research.get("claims") or []
    sources = research.get("sources") or []
    questions = research.get("remaining_questions") or []
    return {
        "title": title,
        "summary": f"Evidence report on {brief}. Research mode: {research.get('mode', 'rapid')}.",
        "sections": [
            {
                "heading": "Synthesis",
                "paragraphs": [str(research.get("message") or "No synthesis was returned.")[:12000]],
                "bullets": [],
            },
            {
                "heading": "Evidence Claims",
                "paragraphs": [],
                "bullets": [
                    f"{claim.get('claim_id', 'Claim')}: {claim.get('text') or claim.get('claim') or ''} "
                    f"[{', '.join(claim.get('source_ids') or [])}] ({claim.get('status', 'uncertain')})"
                    for claim in claims[:30]
                ],
            },
            {
                "heading": "Sources",
                "paragraphs": [],
                "bullets": [
                    f"[{source.get('source_id', 'Source')}] {source.get('title', 'Untitled')} - {source.get('url', '')}"
                    for source in sources[:30]
                ],
            },
            {
                "heading": "Remaining Questions",
                "paragraphs": [],
                "bullets": [str(question)[:1000] for question in questions[:20]],
            },
        ],
    }


def _safe_sections(value: list[Any]) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    for item in value[:20]:
        if not isinstance(item, dict):
            continue
        heading = str(item.get("heading") or "").strip()[:200]
        paragraphs = [str(text).strip()[:5000] for text in item.get("paragraphs", [])[:20] if str(text).strip()]
        bullets = [str(text).strip()[:1000] for text in item.get("bullets", [])[:30] if str(text).strip()]
        if heading and (paragraphs or bullets):
            sections.append({"heading": heading, "paragraphs": paragraphs, "bullets": bullets})
    return sections


def _verify_docx(path: Path) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            package_ok = {"[Content_Types].xml", "word/document.xml"} <= names
        checks.append(
            {"name": "docx_package", "passed": package_ok, "evidence": "required Word parts present" if package_ok else "required Word parts missing"}
        )
        from docx import Document

        reopened = Document(path)
        paragraphs = [paragraph.text for paragraph in reopened.paragraphs if paragraph.text.strip()]
        content_ok = len(paragraphs) >= 2
        checks.append(
            {"name": "docx_content", "passed": content_ok, "evidence": f"{len(paragraphs)} non-empty paragraphs reopened"}
        )
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as error:
        return {
            "ok": False,
            "paragraphs": 0,
            "checks": [{"name": "docx_reopen", "passed": False, "evidence": str(error)[:300]}],
            "verification": "The Word document could not be reopened and was not accepted as verified.",
        }
    ok = all(check["passed"] for check in checks)
    return {
        "ok": ok,
        "paragraphs": len(paragraphs),
        "checks": checks,
        "verification": (
            f"Verified: the DOCX package reopened successfully with {len(paragraphs)} non-empty paragraphs; size and SHA-256 were recorded."
            if ok
            else "The DOCX package exists, but its required content checks did not pass."
        ),
    }


def _prepare_work_root(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    if root.is_symlink():
        raise RuntimeError("The AURIS managed work root cannot be a symbolic link.")
    return root.resolve(strict=True)


def _file_artifact(
    kind: str,
    title: str,
    path: Path,
    format_name: str,
    generator_info: dict[str, Any],
    checks: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "kind": kind,
        "name": title,
        "path": str(path),
        "format": format_name,
        "bytes": path.stat().st_size,
        "sha256": _sha256_file(path),
        "generator": generator_info,
        "checks": checks,
    }


def _create_unique_directory(parent: Path, slug: str) -> Path:
    parent.mkdir(parents=True, exist_ok=True)
    if parent.is_symlink():
        raise RuntimeError("The AURIS project directory cannot be a symbolic link.")
    for index in range(1, 1000):
        name = slug if index == 1 else f"{slug}-{index}"
        candidate = parent / name
        try:
            candidate.mkdir()
        except FileExistsError:
            continue
        return candidate
    raise RuntimeError("A unique managed project directory could not be allocated.")


def _unique_path(parent: Path, slug: str, suffix: str) -> Path:
    for index in range(1, 1000):
        name = f"{slug}{suffix}" if index == 1 else f"{slug}-{index}{suffix}"
        candidate = parent / name
        if not candidate.exists():
            return candidate
    raise RuntimeError("A unique managed document path could not be allocated.")


def _json_object(text: str) -> dict[str, Any]:
    clean = text.strip()
    clean = re.sub(r"^```(?:json)?\s*", "", clean, flags=re.IGNORECASE)
    clean = re.sub(r"\s*```$", "", clean)
    start = clean.find("{")
    end = clean.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("No JSON object was returned.")
    payload = json.loads(clean[start : end + 1])
    if not isinstance(payload, dict):
        raise ValueError("The generated payload is not an object.")
    return payload


def _requested_name(command: str) -> str:
    quoted = re.search(r"\b(?:called|named|titled)\s+[\"']([^\"']+)[\"']", command, flags=re.IGNORECASE)
    if quoted:
        return quoted.group(1).strip()[:100]
    match = re.search(r"\b(?:called|named|titled)\s+[\"']?([^\"']+?)[\"']?(?:\s+(?:about|on|for|that|to)|$)", command, flags=re.IGNORECASE)
    return match.group(1).strip(" .")[:100] if match else ""


def _title_from_brief(brief: str, fallback: str) -> str:
    words = re.findall(r"[A-Za-z0-9]+", brief)
    if not words:
        return fallback
    ignored = {"a", "an", "the", "about", "on", "for", "of", "to", "my", "me", "please"}
    selected = [word for word in words if word.casefold() not in ignored][:8]
    return " ".join(selected).title() or fallback


def _slug(value: str, fallback: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-").casefold()[:60]
    return slug or fallback


def _clean_command(command: str) -> str:
    clean = command.strip().strip(" .")
    clean = re.sub(r"^(?:hey\s+)?auris(?:\s*,\s*|\s+)", "", clean, flags=re.IGNORECASE)
    clean = re.sub(
        r"^(?:please\s+|(?:(?:can|could|would|will)\s+you\s+)(?:please\s+)?|i\s+want\s+you\s+to\s+)",
        "",
        clean,
        flags=re.IGNORECASE,
    )
    return clean.strip(" .")


def _html_text(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _xml_text(value: Any) -> str:
    return _html_text(str(value)).replace("\n", "<br/>")


def _safe_spreadsheet_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value).strip()[:2000]
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(64 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
